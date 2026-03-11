from pathlib import Path

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QListWidget, QListWidgetItem, QMessageBox
)
from PyQt6.QtCore import Qt, pyqtSignal
from PyQt6.QtGui import QDragEnterEvent, QDropEvent

from app.modules.db.storage import DocumentDB


def _extract_text(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".txt":
        return p.read_text(encoding="utf-8", errors="ignore")

    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())

    elif suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join(str(c) if c is not None else "" for c in row)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)

    elif suffix == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            raise ValueError("PDF 지원을 위해 pypdf를 설치하세요: pip install pypdf")

    else:
        raise ValueError(f"지원하지 않는 파일 형식: {suffix}\n지원 형식: TXT, DOCX, XLSX, PDF")


class DropArea(QLabel):
    files_dropped = pyqtSignal(list)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setText("파일을 여기에 드래그하여 추가\nTXT · DOCX · XLSX · PDF")
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setAcceptDrops(True)
        self.setMinimumHeight(120)
        self._set_normal_style()

    def _set_normal_style(self):
        self.setStyleSheet("""
            QLabel {
                border: 2px dashed #5A5A7A;
                border-radius: 8px;
                color: #9090A8;
                font-size: 10pt;
                padding: 16px;
            }
        """)

    def _set_hover_style(self):
        self.setStyleSheet("""
            QLabel {
                border: 2px dashed #7878FF;
                border-radius: 8px;
                color: #C0C0D8;
                font-size: 10pt;
                padding: 16px;
                background-color: #1E1E32;
            }
        """)

    def dragEnterEvent(self, event: QDragEnterEvent):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self._set_hover_style()
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self._set_normal_style()

    def dropEvent(self, event: QDropEvent):
        self._set_normal_style()
        paths = [url.toLocalFile() for url in event.mimeData().urls()]
        if paths:
            self.files_dropped.emit(paths)


class DBWidget(QWidget):
    status_message = pyqtSignal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self._build_ui()
        self._refresh_list()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(32, 28, 32, 28)
        layout.setSpacing(16)

        # 헤더
        title = QLabel("DB 자료 저장")
        title.setObjectName("pageTitle")
        layout.addWidget(title)

        desc = QLabel("파일을 드래그하면 저장되며, 보고서·제안서 AI 생성 시 참고 자료로 자동 활용됩니다.")
        desc.setStyleSheet("color: #9090A8; font-size: 9pt;")
        desc.setWordWrap(True)
        layout.addWidget(desc)

        # 드래그 영역
        self.drop_area = DropArea()
        self.drop_area.files_dropped.connect(self._on_files_dropped)
        layout.addWidget(self.drop_area)

        # 목록 헤더
        list_header = QHBoxLayout()
        lbl = QLabel("저장된 자료")
        lbl.setObjectName("sectionLabel")
        list_header.addWidget(lbl)
        list_header.addStretch()
        self.count_label = QLabel("0개")
        self.count_label.setStyleSheet("color: #9090A8; font-size: 9pt;")
        list_header.addWidget(self.count_label)
        layout.addLayout(list_header)

        # 목록
        self.list_widget = QListWidget()
        self.list_widget.setAlternatingRowColors(True)
        layout.addWidget(self.list_widget, 1)

        # 삭제 버튼
        btn_row = QHBoxLayout()
        self.btn_delete = QPushButton("선택 항목 삭제")
        self.btn_delete.setObjectName("secondaryButton")
        self.btn_delete.clicked.connect(self._delete_selected)
        btn_row.addWidget(self.btn_delete)
        btn_row.addStretch()
        layout.addLayout(btn_row)

    def _refresh_list(self):
        self.list_widget.clear()
        docs = DocumentDB.get_all()
        for doc in docs:
            item = QListWidgetItem(f"  📄  {doc['name']}    ({doc['added_date']})")
            item.setData(Qt.ItemDataRole.UserRole, doc["id"])
            self.list_widget.addItem(item)
        self.count_label.setText(f"{len(docs)}개")

    def _on_files_dropped(self, paths: list[str]):
        added, errors = 0, []
        for path in paths:
            try:
                content = _extract_text(path)
                DocumentDB.add(Path(path).name, content)
                added += 1
            except Exception as e:
                errors.append(f"{Path(path).name}: {e}")

        if added:
            self._refresh_list()
            self.status_message.emit(f"{added}개 파일이 DB에 추가되었습니다.")

        if errors:
            QMessageBox.warning(self, "일부 파일 오류", "\n".join(errors))

    def _delete_selected(self):
        selected = self.list_widget.selectedItems()
        if not selected:
            QMessageBox.information(self, "안내", "삭제할 항목을 선택하세요.")
            return
        reply = QMessageBox.question(self, "삭제 확인", f"{len(selected)}개 항목을 삭제하시겠습니까?")
        if reply == QMessageBox.StandardButton.Yes:
            for item in selected:
                DocumentDB.remove(item.data(Qt.ItemDataRole.UserRole))
            self._refresh_list()
            self.status_message.emit("선택 항목이 삭제되었습니다.")
